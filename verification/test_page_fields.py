"""Local-only checks for the narrow DOCX page-field inspection exception."""
import hashlib
import unicodedata
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from backend import docx_engine
from backend.docx_io import W_NS, R_NS, _find_uninspected, _is_simple_page_field, read_package

W = '{' + W_NS + '}'


def field(instruction='PAGE', content='<w:r><w:t>1</w:t></w:r>'):
    result = etree.fromstring(
        f'<w:fldSimple xmlns:w="{W_NS}" xmlns:r="{R_NS}">{content}</w:fldSimple>'.encode()
    )
    result.set(W + 'instr', instruction)
    return result


def labels(element):
    return _find_uninspected(etree.tostring(element))


@pytest.mark.parametrize('instruction', [
    'PAGE', 'NUMPAGES', 'SECTIONPAGES', ' page ',
    ' PAGE \\* MERGEFORMAT ', 'NUMPAGES\t\\*\tMERGEFORMAT',
    'SECTIONPAGES \\* mergeformat',
])
@pytest.mark.parametrize('content', [
    '', '<w:r/>', '<w:r><w:t/></w:r>',
    '<w:r><w:t>1</w:t></w:r>',
    '<w:r><w:t>12</w:t><w:t>3</w:t></w:r><w:r><w:t>4</w:t></w:r>',
    '<w:r w:rsidR="00ABCDEF"><w:t xml:space="preserve">12</w:t></w:r>',
])
def test_only_known_page_instruction_and_empty_or_ascii_numeric_cache(instruction, content):
    element = field(instruction, content)
    before = etree.tostring(element)
    assert _is_simple_page_field(element)
    assert labels(element) == []
    assert etree.tostring(element) == before


@pytest.mark.parametrize('instruction', [
    '', 'AUTHOR', 'FILENAME', 'DATE', 'DOCVARIABLE PAGE',
    'PAGE AUTHOR', 'PAGE \\* MERGEFORMAT AUTHOR',
    'PAGE \\* MERGEFORMAT \\* MERGEFORMAT',
    'PAGE \\* ROMAN', 'NUMPAGES \\# "0"', 'PAGE+1',
    'INCLUDETEXT "https://example.test/a"',
    'PAGE\u00a0\\* MERGEFORMAT',
])
def test_unknown_or_compound_instructions_remain_uninspected(instruction):
    assert labels(field(instruction)) == ['필드']


@pytest.mark.parametrize('content', [
    '<w:r><w:t>홍길동</w:t></w:r>',
    '<w:r><w:t>1 example@example.test</w:t></w:r>',
    '<w:r><w:t>１</w:t></w:r>', '<w:r><w:t>Ⅲ</w:t></w:r>',
    '<w:r><w:t>-1</w:t></w:r>', '<w:r><w:t>1 2</w:t></w:r>',
    '<w:r><w:t> 1 </w:t></w:r>',
    '<w:r><w:instrText>AUTHOR</w:instrText><w:t>1</w:t></w:r>',
    '<w:r><w:fldChar w:fldCharType="begin"/><w:t>1</w:t></w:r>',
    '<w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>',
    '<w:hyperlink r:id="rId9"><w:r><w:t>1</w:t></w:r></w:hyperlink>',
    '<w:r><w:drawing/><w:t>1</w:t></w:r>',
    '<w:r><w:t><w:externalLink/></w:t></w:r>',
    '<w:r><w:tab/><w:t>1</w:t></w:r>',
    '<w:r><w:rPr><w:b/></w:rPr><w:t>1</w:t></w:r>',
    '<w:r><w:t r:id="rId9">1</w:t></w:r>',
    '<w:r r:id="rId9"><w:t>1</w:t></w:r>',
    '<w:r><w:t>1</w:t>hidden result</w:r>',
    'hidden result<w:r><w:t>1</w:t></w:r>',
    '<w:r><!--hidden--><w:t>1</w:t></w:r>',
    '<x:r xmlns:x="urn:unknown"><x:t>1</x:t></x:r>',
])
def test_unsafe_or_unrecognized_result_structure_remains_uninspected(content):
    assert '필드' in labels(field(content=content))


def test_attributes_namespace_and_other_field_warnings_remain_checked():
    valid = field()
    valid.set(W + 'dirty', 'true')
    valid.set(W + 'fldLock', '0')
    assert labels(valid) == []
    valid.set('{' + R_NS + '}id', 'rId9')
    assert '필드' in labels(valid)
    wrong_namespace = etree.fromstring(b'<fldSimple instr="PAGE"><r><t>1</t></r></fldSimple>')
    assert '필드' in labels(wrong_namespace)
    mixed = etree.Element(W + 'p', nsmap={'w': W_NS})
    mixed.append(field())
    mixed.append(field('AUTHOR'))
    assert labels(mixed) == ['필드']
    only_complex_marker = etree.Element(W + 'fldChar', nsmap={'w': W_NS})
    only_complex_marker.set(W + 'fldCharType', 'begin')
    assert labels(only_complex_marker) == ['필드']


def test_inspect_and_saved_copy_preserve_page_field_and_original(tmp_path):
    source, target = tmp_path/'original.docx', tmp_path/'copy.docx'
    document = Document()
    document.add_paragraph('합성 테스트 문서')
    document.sections[0].footer.paragraphs[0]._p.append(field('PAGE \\* MERGEFORMAT'))
    document.save(source)
    original = source.read_bytes()
    original_footer = read_package(source)['members']['word/footer1.xml']
    inspection = docx_engine.inspect(source)
    assert not any('필드' in label for label in inspection['uninspected'])
    assert any(u['text'] == '1' for u in inspection['units'])
    metadata_actions = {item['id']: 'keep' for item in inspection['metadata']}
    result = docx_engine.render(source, target, inspection, [], metadata_actions)
    assert result['passed']
    assert source.read_bytes() == original
    assert read_package(target)['members']['word/footer1.xml'] == original_footer
    assert not any('필드' in label for label in docx_engine.inspect(target)['uninspected'])


def test_provided_document_page_field_local_only():
    candidates = [p for p in Path('/Users/a1717771/Downloads').glob('*.docx')
                  if unicodedata.normalize('NFC', p.name) == '가리미_테스트_출장등록및정산.docx']
    if not candidates:
        pytest.skip('사용자 제공 문서는 공개 테스트 자료에 복사하지 않습니다.')
    path = candidates[0]
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    package = read_package(path)
    footer = etree.fromstring(package['members']['word/footer1.xml'])
    fields = list(footer.iter(W + 'fldSimple'))
    assert len(fields) == 1
    assert fields[0].get(W + 'instr') == 'PAGE'
    assert _is_simple_page_field(fields[0])
    inspection = docx_engine.inspect(path)
    assert not any('필드' in label for label in inspection['uninspected'])
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
